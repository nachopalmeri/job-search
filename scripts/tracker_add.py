"""Agrega y actualiza postulaciones en el Excel de Nacho (Postulaciones_Ignacio_Palmeri.xlsx).

El Excel es el registro oficial (archivo de Nacho, no depende de ninguna suscripcion).
Cada cambio deja ademas un espejo en tracker/postulaciones.csv y un backup del Excel.

Uso (Python del venv, que tiene openpyxl):
  D:\\Dev\\wenv\\Scripts\\python.exe scripts\\tracker_add.py list [--estado Postulada]
  D:\\Dev\\wenv\\Scripts\\python.exe scripts\\tracker_add.py followups
  D:\\Dev\\wenv\\Scripts\\python.exe scripts\\tracker_add.py add --empresa X --puesto Y --link URL --fuente LinkedIn --cv archivo.pdf
  D:\\Dev\\wenv\\Scripts\\python.exe scripts\\tracker_add.py update 12 --estado Entrevista --log "Me escribio la recruiter"
  D:\\Dev\\wenv\\Scripts\\python.exe scripts\\tracker_add.py export-csv
  D:\\Dev\\wenv\\Scripts\\python.exe scripts\\tracker_add.py export-obsidian   (regenera el registro en el vault)

Con el Excel abierto en Excel el archivo esta bloqueado: el script avisa y no escribe.
Cada add/update/export-csv regenera tambien el registro en Obsidian (tabla + una nota por postulacion).
Carpeta: --obsidian-dir, o la variable OBSIDIAN_DIR, o la ruta por defecto de abajo; --no-obsidian lo omite.
Ruta del Excel: --xlsx, o la variable TRACKER_XLSX, o la ruta por defecto de abajo.
"""
import argparse, csv, datetime as dt, os, re, shutil, sys, unicodedata
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("Falta openpyxl. Usa el Python del venv: D:\\Dev\\wenv\\Scripts\\python.exe scripts\\tracker_add.py ...")

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = r"C:\Users\Nacho\Documents\Busqueda laboral\Postulaciones_Ignacio_Palmeri.xlsx"
CSV_PATH = ROOT / "tracker" / "postulaciones.csv"
DEFAULT_OBSIDIAN = r"D:\Facultad\UADE-Vault\Efforts\Trabajo\Postulaciones"
AUTO_START, AUTO_END = "<!-- job-search:auto -->", "<!-- /job-search:auto -->"
SHEET, HDR_ROW, FIRST_ROW = "Postulaciones", 3, 4

# clave interna -> encabezado de la hoja Postulaciones (el orden del CSV espejo es el de esta lista)
FIELDS = [
    ("n", "N°"), ("empresa", "Empresa"), ("puesto", "Puesto"), ("estado", "Estado"),
    ("fecha_postulacion", "Fecha postulación"), ("ultima_interaccion", "Última interacción"),
    ("proxima_accion", "Próxima acción"), ("link", "Link del aviso"), ("fuente", "Dónde la encontré"),
    ("canal", "Canal"), ("referido", "Recomendado por / networking"), ("descripcion", "Descripción del puesto"),
    ("cv", "CV usado"), ("cv_ruta", "Ruta / link del CV"), ("carta", "Mensaje o carta"),
    ("modalidad", "Modalidad"), ("ubicacion", "Ubicación"), ("jornada", "Jornada / horario"),
    ("remuneracion", "Remuneración"), ("fit", "Fit (1-5)"), ("prioridad", "Prioridad"), ("area", "Área"),
    ("fecha_limite", "Fecha límite"), ("contacto_nombre", "Contacto"), ("contacto_cargo", "Cargo del contacto"),
    ("contacto_email", "Email / LinkedIn contacto"), ("etapa", "Etapa / resultado"), ("historial", "Historial"),
    ("notas", "Notas"), ("motivo_cierre", "Motivo de cierre"),
]
DATE_KEYS = {"fecha_postulacion", "ultima_interaccion", "fecha_limite"}
ESTADOS = ["Para postular", "Postulada", "Contactado", "Entrevista", "Prueba técnica", "Oferta", "Rechazada", "Descartada", "Retirada"]
ACTIVOS = {"Postulada", "Contactado"}
EN_CURSO = {"Entrevista", "Prueba técnica"}
CERRADOS = {"Oferta", "Rechazada", "Descartada", "Retirada"}
DIAS_FOLLOWUP, DIAS_FRIO = 3, 14


def _n(v):
    """N° como entero; 0 si la celda está vacía o no es número (fila cargada a mano)."""
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return 0


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str) and v.strip():
        for f in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
            try:
                return dt.datetime.strptime(v.strip(), f).date()
            except ValueError:
                pass
    return None


class Tracker:
    def __init__(self, path):
        self.path = Path(path)
        if not self.path.exists():
            sys.exit(f"No existe el Excel: {self.path}\nPasa --xlsx o definí TRACKER_XLSX.")
        self.wb = load_workbook(self.path)
        self.ws = self.wb[SHEET]
        head = {str(self.ws.cell(HDR_ROW, c).value or "").strip(): c for c in range(1, self.ws.max_column + 1)}
        missing = [h for _, h in FIELDS if h not in head]
        if missing:
            sys.exit(f"Al Excel le faltan columnas (¿lo reordenaste?): {missing}")
        self.col = {k: head[h] for k, h in FIELDS}

    def locked(self):
        return (self.path.parent / f"~${self.path.name}").exists()

    def rows(self):
        out = []
        for r in range(FIRST_ROW, self.ws.max_row + 1):
            if self.ws.cell(r, self.col["empresa"]).value:
                out.append((r, {k: self.ws.cell(r, c).value for k, c in self.col.items()}))
        return out

    def find(self, ident, empresa=None, puesto=None):
        rs = self.rows()
        if ident:
            for r, d in rs:
                if str(d["n"]).strip() in (str(ident), f"{ident}.0"):
                    return r, d
            sys.exit(f"No hay ninguna postulación con N° {ident}.")
        hits = [(r, d) for r, d in rs if norm(empresa) in norm(d["empresa"]) and (not puesto or norm(puesto) in norm(d["puesto"]))]
        if len(hits) != 1:
            sys.exit(f"{len(hits)} coincidencias para empresa={empresa!r} puesto={puesto!r}: " + "; ".join(f"#{d['n']} {d['empresa']} · {d['puesto']}" for _, d in hits[:6]))
        return hits[0]

    def set(self, r, key, val):
        cell = self.ws.cell(r, self.col[key])
        if key in DATE_KEYS:
            val = to_date(val) if val else None
            cell.number_format = "dd/mm/yyyy"
        cell.value = val if val not in ("", None) else None

    def save(self):
        if self.locked():
            sys.exit("El Excel está abierto (existe el archivo de bloqueo ~$...). Cerralo en Excel y repetí el comando.")
        backup_dir = self.path.parent / "backup"
        backup_dir.mkdir(exist_ok=True)
        shutil.copy2(self.path, backup_dir / (self.path.stem + ".anterior.xlsx"))
        try:
            self.wb.save(self.path)
        except PermissionError:
            sys.exit("No pude escribir el Excel (está abierto o bloqueado). Cerralo y repetí.")

    def export_csv(self):
        CSV_PATH.parent.mkdir(exist_ok=True)
        keys = [k for k, _ in FIELDS]
        with open(CSV_PATH, "w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=keys)
            w.writeheader()
            for _, d in self.rows():
                row = {}
                for k in keys:
                    v = d[k]
                    if k in DATE_KEYS:
                        v = (to_date(v) or "")
                        v = v.isoformat() if v else ""
                    elif k == "historial":  # dd/mm/aa · texto  ->  AAAA-MM-DD · texto
                        lines = []
                        for ln in str(v or "").splitlines():
                            m = re.match(r"(\d{2}/\d{2}/\d{2,4}) · (.*)", ln)
                            dd = to_date(m.group(1)) if m else None
                            lines.append(f"{dd.isoformat()} · {m.group(2)}" if dd else ln)
                        v = "\n".join(lines)
                    elif k == "n":
                        v = _n(v) if v not in (None, "") else ""
                    row[k] = "" if v is None else v
                w.writerow(row)
        return CSV_PATH


def _md(v):
    return str(v).replace("|", "\\|").replace("\n", " ").strip() if v not in (None, "") else ""


def _yq(v):
    """Valor seguro dentro de comillas dobles de YAML (frontmatter)."""
    return str(v if v is not None else "").replace("\\", "/").replace('"', "'").replace("\n", " ").strip()


def _fname(d):
    base = f"{_n(d['n']):03d} {d['empresa']} - {d['puesto']}"
    return re.sub(r"\s+", " ", re.sub(r'[\\/:*?"<>|#^\[\]]', " ", base)).strip()[:110]


def _file_link(path):
    p = str(path or "").strip()
    if re.match(r"^[A-Za-z]:[\\/]", p):
        return f"[{Path(p).name}](<file:///{p.replace(chr(92), '/')}>)"
    return p


def _fit(v):
    try:
        n = round(float(v))
        return f"{'★' * n}{'☆' * (5 - n)} ({float(v):g}/5)"
    except (TypeError, ValueError):
        return ""


def export_obsidian(t, folder):
    """Tabla resumen + una nota por postulacion. Lo generado va entre marcadores; lo que Nacho escriba fuera de ellos se conserva."""
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    hoy = dt.date.today()
    rows = sorted(t.rows(), key=lambda x: -_n(x[1]["n"] or 0))
    lines = ["---", "tipo: registro", "area: trabajo", f"actualizado: {hoy.isoformat()}",
             "fuente: Excel Postulaciones_Ignacio_Palmeri.xlsx (generado por scripts/tracker_add.py)", "---", "",
             "# Registro de postulaciones", "",
             f"Generado automáticamente desde el Excel ({len(rows)} postulaciones, {hoy.strftime('%d/%m/%Y')}). No lo edites a mano: editá el Excel o la nota de cada postulación (debajo de la marca).", "",
             "**Por estado:** " + " · ".join(f"{e} {n}" for e, n in ((e, sum(1 for _, d in rows if d["estado"] == e)) for e in ESTADOS) if n)
             + " — **Por fuente:** " + " · ".join(f"{f} {n}" for f, n in sorted(((f, sum(1 for _, d in rows if (d["fuente"] or "?") == f)) for f in {(d["fuente"] or "?") for _, d in rows}), key=lambda x: -x[1])[:6]), "",
             "| N° | Empresa | Puesto | Estado | Seguimiento | Postulé | Match | CV usado | Ruta del CV | Aviso | Dónde | Canal |",
             "|---|---|---|---|---|---|---|---|---|---|---|---|"]
    for _, d in rows:
        f = to_date(d["fecha_postulacion"])
        note = _fname(d)
        link = f"[aviso]({d['link']})" if str(d["link"] or "").startswith("http") else _md(d["link"])
        lines.append("| " + " | ".join([str(_n(d["n"])), _md(d["empresa"]), f"[[{note}\\|{_md(d['puesto'])}]]", _md(d["estado"]), seguimiento(d, hoy),
                                        f.strftime("%d/%m/%y") if f else "", _fit(d["fit"]), _md(d["cv"]), _file_link(d["cv_ruta"]), link, _md(d["fuente"]), _md(d["canal"])]) + " |")
        props = ["---", "tipo: postulacion", f"n: {_n(d['n'])}", f"empresa: \"{_yq(d['empresa'])}\"", f"puesto: \"{_yq(d['puesto'])}\"",
                 f"estado: {_yq(d['estado'])}", f"fecha_postulacion: {(f.isoformat() if f else '')}", f"fit: {d['fit'] if d['fit'] not in (None, '') else ''}",
                 f"link: \"{_yq(d['link'])}\"", f"cv: \"{_yq(d['cv'])}\"", f"cv_ruta: \"{_yq(d['cv_ruta'])}\"", "---"]
        lim = to_date(d["fecha_limite"])
        body = [AUTO_START, f"# {d['empresa']} · {d['puesto']}", "",
                f"- **Estado:** {d['estado']} · **Seguimiento:** {seguimiento(d, hoy)}",
                f"- **Aviso:** {d['link'] or '-'}",
                f"- **Dónde la encontré:** {d['fuente'] or '-'} · **Canal:** {d['canal'] or '-'}" + (f" · **Recomendado por:** {d['referido']}" if d["referido"] else ""),
                f"- **Match (fit):** {_fit(d['fit']) or '-'} · **Prioridad:** {d['prioridad'] or '-'}",
                f"- **CV usado:** {d['cv'] or '-'} → {_file_link(d['cv_ruta']) or '-'}",
                f"- **Modalidad / ubicación / jornada:** {d['modalidad'] or '-'} / {d['ubicacion'] or '-'} / {d['jornada'] or '-'} · **Remuneración:** {d['remuneracion'] or '-'}",
                f"- **Área:** {d['area'] or '-'} · **Fecha límite:** {lim.strftime('%d/%m/%Y') if lim else '-'}",
                f"- **Contacto:** {' · '.join(str(x) for x in (d['contacto_nombre'], d['contacto_cargo'], d['contacto_email']) if x) or '-'}",
                f"- **Etapa / resultado:** {d['etapa'] or '-'} · **Próxima acción:** {d['proxima_accion'] or '-'}"]
        for title, key in (("Descripción del puesto", "descripcion"), ("Mensaje o carta", "carta"), ("Historial", "historial"), ("Notas", "notas"), ("Motivo de cierre", "motivo_cierre")):
            if d[key]:
                body += ["", f"## {title}", "", str(d[key]).strip()]
        body.append(AUTO_END)
        path = folder / f"{note}.md"
        prev = [q for q in folder.glob(f"{_n(d['n']):03d} *.md") if q != path]
        if prev and not path.exists():  # cambió empresa/puesto: la nota vieja se renombra y conserva "Mis apuntes"
            prev[0].rename(path)
        keep = "\n\n## Mis apuntes\n\n"
        if path.exists():
            old = path.read_text(encoding="utf-8")
            if AUTO_END not in old:  # nota escrita a mano o sin marca: no se pisa
                print(f"AVISO: {path.name} existe sin la marca {AUTO_END}; no la toqué.")
                continue
            tail = old.split(AUTO_END, 1)[1]
            if tail.strip():
                keep = tail.rstrip("\n") + "\n"
        path.write_text("\n".join(props + [""] + body) + keep, encoding="utf-8", newline="\n")
    (folder / "Registro de postulaciones.md").write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")
    return folder


def seguimiento(d, hoy):
    """Misma regla que la columna 'Seguimiento' del Excel."""
    est = d["estado"]
    lim, ult = to_date(d["fecha_limite"]), max(filter(None, [to_date(d["fecha_postulacion"]), to_date(d["ultima_interaccion"])]), default=None)
    if est == "Para postular":
        if lim and lim < hoy:
            return "Aviso vencido"
        return f"Vence {lim.strftime('%d/%m')}" if lim and (lim - hoy).days <= 7 else "Pendiente de enviar"
    if est in CERRADOS:
        return "Cerrada"
    if not ult:
        return "Falta fecha"
    n = (hoy - ult).days
    if est in EN_CURSO:
        return f"Revisar ({n} d)" if n >= DIAS_FOLLOWUP else "En curso"
    if n >= DIAS_FRIO:
        return f"Sin movimiento ({n} d)"
    return f"Hacer follow-up ({n} d)" if n >= DIAS_FOLLOWUP else "Esperar"


def cmd_list(a, t):
    hoy = dt.date.today()
    rows = [(r, d) for r, d in t.rows() if not a.estado or d["estado"] == a.estado]
    for _, d in rows:
        f = to_date(d["fecha_postulacion"])
        print(f"#{_n(d['n']):>3} {d['estado']:<13} {f.strftime('%d/%m') if f else '   -  '}  {str(d['empresa'])[:26]:<26} | {str(d['puesto'])[:52]:<52} | {seguimiento(d, hoy)}")
    print(f"\n{len(rows)} postulaciones ({t.path.name})")


def cmd_followups(a, t):
    hoy = dt.date.today()
    buckets = {"follow": [], "revisar": [], "frio": [], "pendiente": [], "vencido": []}
    for _, d in t.rows():
        s = seguimiento(d, hoy)
        line = f"#{_n(d['n'])} {d['empresa']} · {d['puesto']} ({s})"
        if d["canal"] in ("Recomendación", "Networking") and d["referido"]:
            line += f"\n     ↳ llegó por {d['canal'].lower()}: el mensaje va a {d['referido']}, no a la empresa"
        if s.startswith(("Hacer follow-up", "Revisar")) and d["proxima_accion"]:
            line += f"\n     ↳ próxima acción anotada: {str(d['proxima_accion'])[:110]}"
        if s.startswith(("Hacer follow-up", "Revisar")) and d["notas"]:
            line += f"\n     ↳ notas: {str(d['notas'])[:110]}"
        for key, pre in (("follow", "Hacer follow-up"), ("revisar", "Revisar"), ("frio", "Sin movimiento"), ("vencido", "Aviso vencido")):
            if s.startswith(pre):
                buckets[key].append(line)
        if s.startswith("Vence") or s == "Pendiente de enviar":
            buckets["pendiente"].append(line)
    titles = {"follow": "Hacer follow-up hoy", "revisar": "Revisar (entrevista o prueba en curso)", "frio": "Sin movimiento (decidir: cerrar o último mensaje)",
              "pendiente": "Para postular", "vencido": "Avisos vencidos (marcar Retirada)"}
    for k, title in titles.items():
        if buckets[k]:
            print(f"\n## {title} ({len(buckets[k])})")
            for ln in buckets[k]:
                print(" -", ln)
    print(f"\nHoy: {hoy.strftime('%d/%m/%Y')}. Regla: follow-up a los {DIAS_FOLLOWUP} días, 'sin movimiento' a los {DIAS_FRIO}.")


def cmd_add(a, t):
    if a.estado not in ESTADOS:
        sys.exit(f"Estado inválido {a.estado!r}. Opciones: {ESTADOS}")
    for _, d in t.rows():
        same = norm(d["empresa"]) == norm(a.empresa) and (norm(d["puesto"]) == norm(a.puesto) or (a.link and a.link == d["link"]))
        if same and not a.force:
            sys.exit(f"Ya existe: #{_n(d['n'])} {d['empresa']} · {d['puesto']}. Usá `update` o --force.")
    r = next((r for r in range(FIRST_ROW, t.ws.max_row + 2) if not t.ws.cell(r, t.col["empresa"]).value), None)
    n = max([_n(d["n"]) for _, d in t.rows() if d["n"] not in (None, "")] or [0]) + 1
    hoy = to_date(a.fecha) or dt.date.today()
    vals = dict(n=n, empresa=a.empresa, puesto=a.puesto, estado=a.estado, link=a.link, fuente=a.fuente, canal=a.canal, referido=a.referido,
                descripcion=a.descripcion, cv=a.cv, cv_ruta=a.cv_ruta, carta=a.carta, modalidad=a.modalidad, ubicacion=a.ubicacion, jornada=a.jornada,
                remuneracion=a.remuneracion, fit=a.fit, prioridad=a.prioridad, area=a.area, fecha_limite=a.fecha_limite, contacto_nombre=a.contacto_nombre,
                contacto_cargo=a.contacto_cargo, contacto_email=a.contacto_email, etapa=a.etapa, proxima_accion=a.proxima_accion, notas=a.notas)
    if a.estado != "Para postular":
        vals["fecha_postulacion"] = vals["ultima_interaccion"] = hoy
        vals["historial"] = f"{hoy.strftime('%d/%m/%y')} · {a.log or 'Postulación enviada'}"
    for k, v in vals.items():
        t.set(r, k, v)
    t.save()
    print(f"Agregada #{n} en la fila {r}: {a.empresa} · {a.puesto} ({a.estado})")


def cmd_update(a, t):
    r, d = t.find(a.n, a.empresa, a.puesto)
    hoy = to_date(a.fecha) or dt.date.today()
    changes = {}
    for k in ("estado", "etapa", "proxima_accion", "notas", "motivo_cierre", "link", "cv", "cv_ruta", "carta", "fit", "prioridad",
              "contacto_nombre", "contacto_cargo", "contacto_email", "fecha_limite", "fuente", "canal", "referido"):
        v = getattr(a, k, None)
        if v not in (None, ""):
            changes[k] = v
    if "estado" in changes and changes["estado"] not in ESTADOS:
        sys.exit(f"Estado inválido {changes['estado']!r}. Opciones: {ESTADOS}")
    for pair in a.set or []:
        h, _, v = pair.partition("=")
        key = next((k for k, hh in FIELDS if norm(hh) == norm(h) or k == h), None)
        if not key:
            sys.exit(f"Columna desconocida en --set: {h!r}")
        changes[key] = v
    if changes.get("estado") in ACTIVOS | EN_CURSO and not to_date(d["fecha_postulacion"]):
        changes["fecha_postulacion"] = hoy
    if a.log:
        old = str(d["historial"] or "").strip()
        changes["historial"] = (old + "\n" if old else "") + f"{hoy.strftime('%d/%m/%y')} · {a.log}"
        changes["ultima_interaccion"] = hoy
    elif changes:
        changes.setdefault("ultima_interaccion", hoy)
    if not changes:
        sys.exit("Nada para cambiar. Usá --estado, --log, --etapa, --proxima-accion, --nota, --set 'Columna=valor'…")
    if "notas" in changes and a.append_notas:
        changes["notas"] = (str(d["notas"] or "") + " " + changes["notas"]).strip()
    for k, v in changes.items():
        t.set(r, k, v)
    t.save()
    print(f"Actualizada #{_n(d['n'])} {d['empresa']} · {d['puesto']}: " + ", ".join(f"{k}={str(v)[:40]}" for k, v in changes.items()))


def build_parser():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--xlsx", default=os.environ.get("TRACKER_XLSX", DEFAULT_XLSX))
    p.add_argument("--no-csv", action="store_true", help="no actualizar el espejo tracker/postulaciones.csv")
    p.add_argument("--no-obsidian", action="store_true", help="no regenerar el registro en Obsidian")
    p.add_argument("--obsidian-dir", default=os.environ.get("OBSIDIAN_DIR", DEFAULT_OBSIDIAN))
    sub = p.add_subparsers(dest="cmd", required=True)
    l = sub.add_parser("list"); l.add_argument("--estado")
    sub.add_parser("followups"); sub.add_parser("export-csv"); sub.add_parser("export-obsidian")
    a = sub.add_parser("add")
    a.add_argument("--empresa", required=True); a.add_argument("--puesto", required=True)
    a.add_argument("--estado", default="Postulada"); a.add_argument("--fecha", help="AAAA-MM-DD (default hoy)")
    for k in ("link", "fuente", "referido", "descripcion", "cv", "cv_ruta", "carta", "modalidad", "ubicacion", "jornada", "remuneracion", "area",
              "fecha_limite", "contacto_nombre", "contacto_cargo", "contacto_email", "etapa", "proxima_accion", "notas", "log"):
        a.add_argument("--" + k.replace("_", "-"), dest=k)
    a.add_argument("--canal", default="Postulación directa"); a.add_argument("--prioridad", default="Media")
    a.add_argument("--fit", type=float); a.add_argument("--force", action="store_true")
    u = sub.add_parser("update")
    u.add_argument("n", nargs="?", help="N° de la postulación (o usá --empresa/--puesto)")
    u.add_argument("--empresa"); u.add_argument("--puesto"); u.add_argument("--fecha")
    for k in ("estado", "etapa", "proxima_accion", "notas", "motivo_cierre", "link", "cv", "cv_ruta", "carta", "prioridad", "contacto_nombre",
              "contacto_cargo", "contacto_email", "fecha_limite", "fuente", "canal", "referido", "log"):
        u.add_argument("--" + k.replace("_", "-"), dest=k)
    u.add_argument("--fit", type=float); u.add_argument("--append-notas", action="store_true")
    u.add_argument("--set", action="append", metavar="Columna=valor")
    return p


def main():
    a = build_parser().parse_args()
    t = Tracker(a.xlsx)
    {"list": cmd_list, "followups": cmd_followups, "add": cmd_add, "update": cmd_update, "export-csv": lambda a, t: None, "export-obsidian": lambda a, t: None}[a.cmd](a, t)
    if a.cmd in ("add", "update", "export-csv") and not a.no_csv:
        print("Espejo CSV:", t.export_csv())
    if a.cmd in ("add", "update", "export-csv", "export-obsidian") and not a.no_obsidian:
        try:
            print("Registro Obsidian:", export_obsidian(t, a.obsidian_dir))
        except OSError as e:
            print(f"AVISO: no pude escribir el registro en Obsidian ({e}). El Excel quedó bien.")


if __name__ == "__main__":
    main()
