"""Regenera Ignacio_Palmeri_CV_ES.docx desde cv_master_es.md con texto UTF-8 correcto.

Uso: python regenerate_cv_docx.py
Requiere: python-docx (pip install python-docx).
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

BASE = Path(__file__).resolve().parent.parent
SRC = BASE / "cv_master_es.md"
OUT = BASE / "Ignacio_Palmeri_CV_ES.docx"

REVISAR_RE = re.compile(r"\s*\[REVISAR:[^\]]*\]\s*", re.UNICODE)


def main() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    lines = [l.rstrip("\n") for l in SRC.read_text(encoding="utf-8").splitlines()]
    for line in lines:
        raw = line.strip()
        if not raw or raw == "---" or raw.startswith("# ") or raw.startswith("> "):
            continue
        if raw.startswith("## "):
            p = doc.add_paragraph()
            r = p.add_run(raw[3:].upper())
            r.bold = True
            continue
        text = REVISAR_RE.sub(" ", raw).replace("**", "")
        if raw.startswith("- "):
            doc.add_paragraph(text[2:], style="List Bullet")
        else:
            doc.add_paragraph(text)

    doc.save(OUT)
    print(f"OK -> {OUT}")


if __name__ == "__main__":
    main()
